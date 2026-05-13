import json
import os
import shutil
from datetime import datetime
from functools import wraps
from pathlib import Path

from flask import Flask, Response, redirect, render_template, request, send_file, session, url_for
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
BOND_DIR = DATA_DIR / "bond_picker"
STRATEGY_DIR = DATA_DIR / "strategy_dashboard"
SPREAD_DIR = DATA_DIR / "spread_monitor"
CONFIG_DIR = BASE_DIR / "config"
UPLOADS_DIR = BASE_DIR / "uploads"

BOND_EXCEL = BOND_DIR / "数据.xlsx"
STRATEGY_HTML = STRATEGY_DIR / "信用债策略仪表盘.html"
SPREAD_JS = SPREAD_DIR / "spread_data.js"
SPREAD_JSON = SPREAD_DIR / "spread_data.json"
MAPPING_FILE = CONFIG_DIR / "映射表.xlsx"
SHEET_NAME = "万得"

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-secret-change-before-deploy")

BONDS_CACHE = []
DATA_TIMESTAMP = "尚未加载"


def site_password():
    return os.environ.get("SITE_PASSWORD", "Abcd123%")


def login_required(func):
    @wraps(func)
    def wrapper(*args, **kwargs):
        if not session.get("authenticated"):
            return redirect(url_for("login", next=request.path))
        return func(*args, **kwargs)
    return wrapper



def format_date_only(value):
    if not value:
        return ""
    text = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text[:19], fmt).strftime("%Y - %m - %d")
        except ValueError:
            pass
    if len(text) >= 10:
        head = text[:10].replace("/", "-")
        parts = head.split("-")
        if len(parts) == 3 and all(p.isdigit() for p in parts):
            return f"{int(parts[0]):04d} - {int(parts[1]):02d} - {int(parts[2]):02d}"
    return text

def file_updated(path: Path):
    if not path.exists():
        return "文件不存在"
    return datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y - %m - %d")


def file_size(path: Path):
    if not path.exists():
        return "-"
    size = path.stat().st_size
    if size >= 1024 * 1024:
        return f"{size / 1024 / 1024:.1f} MB"
    if size >= 1024:
        return f"{size / 1024:.1f} KB"
    return f"{size} B"


def read_excel(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET_NAME]

    n1_value = ws.cell(row=1, column=14).value
    data_date = str(n1_value).strip() if n1_value else ""
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()

    bonds = []
    for row in rows:
        if not row or len(row) < 11:
            continue
        code = str(row[0] or "").strip()
        name = str(row[1] or "").strip()
        term = row[2]
        rating = str(row[3] or "").strip()
        issuer = str(row[4] or "").strip()
        ytm = row[5]
        entity = str(row[6] or "").strip()
        ct = str(row[7] or "").strip()
        sub = str(row[8] or "").strip()
        tech = str(row[9] or "").strip()
        ir = str(row[10] or "").strip()
        if not code or not name or term is None or ytm is None:
            continue
        try:
            term = round(float(term), 4)
            ytm = round(float(ytm), 4)
        except (TypeError, ValueError):
            continue
        bonds.append([code, name, term, rating, issuer, ytm, entity, ct, sub, tech, ir])
    return bonds, data_date


def load_bond_data():
    global BONDS_CACHE, DATA_TIMESTAMP
    if not BOND_EXCEL.exists():
        BONDS_CACHE = []
        DATA_TIMESTAMP = "数据文件未找到"
        return
    try:
        bonds, data_date = read_excel(BOND_EXCEL)
        BONDS_CACHE = bonds
        DATA_TIMESTAMP = format_date_only(data_date) or file_updated(BOND_EXCEL)
    except Exception as exc:
        BONDS_CACHE = []
        DATA_TIMESTAMP = f"读取失败: {exc}"


def status_info():
    return {
        "bond_picker": {
            "updated": DATA_TIMESTAMP,
            "total": f"{len(BONDS_CACHE):,}",
        },
        "strategy_dashboard": {
            "updated": file_updated(STRATEGY_HTML),
            "size": file_size(STRATEGY_HTML),
        },
        "spread_monitor": {
            "updated": file_updated(SPREAD_JS),
            "size": file_size(SPREAD_JS),
        },
        "spread_monitor_json": {
            "updated": file_updated(SPREAD_JSON),
            "size": file_size(SPREAD_JSON),
        },
    }


def save_upload(file_storage, destination: Path, allowed_exts):
    filename = file_storage.filename or ""
    ext = Path(filename).suffix.lower()
    if ext not in allowed_exts:
        raise ValueError(f"文件类型不正确，仅支持: {', '.join(sorted(allowed_exts))}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        backup_name = f"{destination.stem}_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}{destination.suffix}"
        shutil.copy2(destination, UPLOADS_DIR / backup_name)
    file_storage.save(destination)


@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("authenticated"):
        return redirect(url_for("home"))
    error = None
    if request.method == "POST":
        if request.form.get("password", "") == site_password():
            session["authenticated"] = True
            session["login_time"] = datetime.now().strftime("%Y-%m-%d %H:%M")
            return redirect(request.args.get("next") or url_for("home"))
        error = "访问密码错误"
    return render_template("login.html", error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
@login_required
def home():
    return render_template("home.html", status=status_info())


@app.route("/bond-picker")
@login_required
def bond_picker():
    return render_template(
        "bond_picker.html",
        bond_data=json.dumps(BONDS_CACHE, ensure_ascii=False),
        timestamp=DATA_TIMESTAMP,
        total=len(BONDS_CACHE),
    )


@app.route("/strategy-dashboard")
@login_required
def strategy_dashboard():
    if not STRATEGY_HTML.exists():
        return Response("策略仪表盘 HTML 文件不存在，请先到后台上传。", status=404, mimetype="text/plain; charset=utf-8")
    html = STRATEGY_HTML.read_text(encoding="utf-8", errors="replace")
    return Response(html, mimetype="text/html; charset=utf-8")


@app.route("/spread-monitor")
@login_required
def spread_monitor():
    return render_template("spread_monitor.html")


@app.route("/data/spread_monitor/spread_data.js")
@login_required
def spread_data_js():
    if not SPREAD_JS.exists():
        return Response("var SPREAD_DATA = {data: [], total_bonds: 0, update_time: '数据文件未找到'};", mimetype="application/javascript; charset=utf-8")
    return send_file(SPREAD_JS, mimetype="application/javascript")


@app.route("/admin", methods=["GET", "POST"])
@login_required
def admin():
    message = None
    error = None
    if request.method == "POST":
        target = request.form.get("target", "")
        file = request.files.get("file")
        if not file or not file.filename:
            error = "请选择要上传的文件"
        else:
            try:
                if target == "bond_picker":
                    save_upload(file, BOND_EXCEL, {".xlsx", ".xls"})
                    load_bond_data()
                    message = f"择券工具 Excel 上传成功，已加载 {len(BONDS_CACHE):,} 条数据。"
                elif target == "strategy_dashboard":
                    save_upload(file, STRATEGY_HTML, {".html", ".htm"})
                    message = "策略仪表盘 HTML 上传成功。"
                elif target == "spread_monitor_js":
                    save_upload(file, SPREAD_JS, {".js"})
                    message = "利差监控 spread_data.js 上传成功。"
                elif target == "spread_monitor_json":
                    save_upload(file, SPREAD_JSON, {".json"})
                    message = "利差监控 spread_data.json 上传成功。"
                else:
                    error = "未知上传目标"
            except Exception as exc:
                error = f"上传失败: {exc}"
    return render_template("admin.html", status=status_info(), message=message, error=error)



@app.route("/api/status")
@login_required
def api_status():
    return status_info()


for directory in [BOND_DIR, STRATEGY_DIR, SPREAD_DIR, CONFIG_DIR, UPLOADS_DIR]:
    directory.mkdir(parents=True, exist_ok=True)

load_bond_data()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5000"))
    debug = os.environ.get("FLASK_DEBUG", "1") == "1"
    app.run(host="0.0.0.0", port=port, debug=debug)